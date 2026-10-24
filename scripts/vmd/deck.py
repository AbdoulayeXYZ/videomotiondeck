"""Import an existing deck (pptx / pdf) so its facts, numbers and pictures feed the film.

Writes <project>/media/deck/ (pictures, slide renders) and <project>/deck_outline.md (text per slide) —
read the outline, then write the story from the user's own facts (never invent new ones)."""
import shutil, subprocess
from pathlib import Path
from .common import log, die


def _pdf(src, out, pages_dir):
    lines = []
    if shutil.which('pdftoppm'):
        subprocess.run(['pdftoppm', '-r', '110', '-png', str(src), str(pages_dir / 'page')], check=False)
    if shutil.which('pdftotext'):
        txt = subprocess.run(['pdftotext', '-layout', str(src), '-'], capture_output=True, text=True).stdout
        for k, page in enumerate(txt.split('\f')):
            if page.strip():
                lines.append(f'\n## Page {k + 1}\n')
                lines += [f'    {ln.rstrip()}' for ln in page.splitlines() if ln.strip()]
    else:
        try:
            from pypdf import PdfReader
            for k, pg in enumerate(PdfReader(str(src)).pages):
                lines.append(f'\n## Page {k + 1}\n'); lines.append(pg.extract_text() or '')
        except Exception:
            lines.append('(no text extractor: install poppler-utils)')
    return lines


def import_deck(src, project_dir):
    src = Path(src); d = Path(project_dir)
    if not src.exists():
        die(f'not found: {src}')
    media = d / 'media' / 'deck'; media.mkdir(parents=True, exist_ok=True)
    out = [f'# Deck outline — {src.name}', '',
           'Use these facts, numbers and pictures. Do not invent new ones; ask the user when something is missing.', '']
    ext = src.suffix.lower()
    if ext == '.pdf':
        out += _pdf(src, out, media)
    elif ext in ('.pptx', '.ppsx'):
        try:
            from pptx import Presentation
        except ImportError:
            die('pip install python-pptx  (needed to read .pptx)')
        prs = Presentation(str(src))
        def walk(shapes):
            for sh in shapes:
                if getattr(sh, 'shape_type', None) == 6 and hasattr(sh, 'shapes'):   # group
                    yield from walk(sh.shapes)
                else:
                    yield sh
        for k, slide in enumerate(prs.slides, 1):
            out.append(f'\n## Slide {k}\n')
            for shp in walk(slide.shapes):
                if shp.has_text_frame and shp.text_frame.text.strip():
                    out.append('- ' + shp.text_frame.text.strip().replace('\n', ' / '))
                if getattr(shp, 'shape_type', None) == 13 or hasattr(shp, 'image'):
                    try:
                        im = shp.image
                        p = media / f'slide{k:02d}_{shp.shape_id}.{im.ext}'
                        p.write_bytes(im.blob)
                        out.append(f'  - picture: media/deck/{p.name} ({shp.width * 96 // 914400}×{shp.height * 96 // 914400}px on slide)')
                    except Exception:
                        pass
                if getattr(shp, 'has_table', False) and shp.has_table:
                    for row in shp.table.rows:
                        out.append('  | ' + ' | '.join(c.text.strip() for c in row.cells) + ' |')
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                out.append('  - notes: ' + slide.notes_slide.notes_text_frame.text.strip().replace('\n', ' / '))
        soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if soffice and shutil.which('pdftoppm'):     # slide renders, handy to reuse visuals / check the brand
            tmp = d / 'build' / 'deck_pdf'; tmp.mkdir(parents=True, exist_ok=True)
            subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(src)], capture_output=True, timeout=300)
            pdfs = list(tmp.glob('*.pdf'))
            if pdfs:
                subprocess.run(['pdftoppm', '-r', '80', '-png', str(pdfs[0]), str(media / 'render')], check=False)
                out.append('\nSlide renders: media/deck/render-*.png')
    else:
        die('supported: .pptx, .pdf')
    (d / 'deck_outline.md').write_text('\n'.join(out) + '\n', encoding='utf-8')
    n = len(list(media.iterdir()))
    log(f'deck imported -> {d / "deck_outline.md"} + {n} file(s) in {media}')
